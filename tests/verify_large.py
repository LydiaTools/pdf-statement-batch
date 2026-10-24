"""Read actual large-browser exports against original transaction truth."""
from pathlib import Path
from collections import Counter
import json, hashlib
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
truth = json.loads((root/'public/samples/expected.json').read_text())['transactions']
workbook = load_workbook(root/'test-results/large-batch.xlsx', read_only=True, data_only=False)
sheet = workbook['Transactions 1']
if sheet.max_row is None:
    sheet.calculate_dimension(force=True)
assert sheet.max_row == 2001
expected_rows = []
for index in range(500):
    for row in truth[(index % 2)*4:(index % 2)*4+4]:
        expected_rows.append([f'batch-{index:04d}.pdf', len(expected_rows) % 4+1, 'USD', 'No', row['date'], row['description'],
                              float(row['debit']) if row['debit'] else '',
                              float(row['credit']) if row['credit'] else '', float(row['balance'])])
checks = 0
for actual, expected in zip(sheet.iter_rows(min_row=2), expected_rows):
    assert [cell.value for cell in actual] == expected, ([c.value for c in actual], expected)
    assert actual[4].data_type == 's'
    for cell in actual[6:]:
        if cell.value != '':
            assert cell.data_type == 'n'
    assert all(c.data_type != 'f' for c in actual)
    checks += len(expected)
accounting = list(workbook['Files'].iter_rows(min_row=2, values_only=True))
assert len(accounting) == 500
assert Counter(row[2] for row in accounting) == {'parsed': 500}
assert len({row[8] for row in accounting}) == 500
cancel = load_workbook(root/'test-results/canceled-batch.xlsx', read_only=True)
statuses = Counter(row[2] for row in cancel['Files'].iter_rows(min_row=2, values_only=True))
assert sum(statuses.values()) == 500 and statuses['canceled'] > 0
mixed = load_workbook(root/'test-results/file-accounting.xlsx', read_only=True)
mixed_rows = list(mixed['Files'].iter_rows(min_row=2, values_only=True))
mixed_statuses = Counter(row[2] for row in mixed_rows)
assert mixed_statuses == {'parsed': 2, 'duplicate': 1, 'failed': 4}
assert mixed.sheetnames == ['Transactions 1', 'Transactions 2', 'Files', 'Read me']
assert mixed['Transactions 2'].max_row == 2
assert list(mixed['Transactions 2'].values)[1] == ('amount-only.pdf', 1, 'EUR', 'No', '10/09/2026', 'SYNTHETIC FEE', 2.5)
unconfirmed = load_workbook(root/'test-results/unconfirmed-currency.xlsx', read_only=True)
assert list(unconfirmed['Transactions 1'].values)[1][2] == 'Unconfirmed'
assert any('password' in str(row[7]).lower() for row in mixed_rows if row[1] == 'locked.pdf')
assert any('text' in str(row[7]).lower() for row in mixed_rows if row[1] == 'scan.pdf')
record = {'files': 500, 'parsed': 500, 'transactionRows': 2000, 'fieldChecks': checks,
          'uniqueInputHashes': 500, 'layouts': 2, 'scope': 'byte-distinct variants of two synthetic PDFs; no bank accuracy claim',
          'canceledWorkbookStatuses': dict(statuses), 'mixedWorkbookStatuses': dict(mixed_statuses),
          'workbookSha256': hashlib.sha256((root/'test-results/large-batch.xlsx').read_bytes()).hexdigest()}
(root/'test-results/large-workbook-verified.json').write_text(json.dumps(record, indent=2)+'\n')
print(json.dumps(record))

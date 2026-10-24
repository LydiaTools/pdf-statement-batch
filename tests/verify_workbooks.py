"""Independent readback of actual browser downloads, using openpyxl."""
from pathlib import Path
from openpyxl import load_workbook
import json

root = Path(__file__).resolve().parents[1]
truth = json.loads((root/'public/samples/expected.json').read_text())['transactions']
checks = []
for name in ['sample-batch.xlsx', 'cached-offline.xlsx', 'edited-zh.xlsx']:
    workbook = load_workbook(root/'test-results'/name, data_only=False)
    assert workbook.sheetnames == ['Transactions 1', 'Files', 'Read me']
    sheet = workbook.worksheets[0]
    assert sheet.max_row == 9
    for index, expected in enumerate(truth, 2):
        got = list(next(sheet.iter_rows(min_row=index, max_row=index, values_only=True)))
        wanted = [expected['file'], ((index-2) % 4)+1, 'USD',
                  'Yes' if name == 'edited-zh.xlsx' and index <= 5 else 'No',
                  expected['date'], expected['description'],
                  float(expected['debit']) if expected['debit'] else '',
                  float(expected['credit']) if expected['credit'] else '',
                  float(expected['balance'])]
        if name == 'edited-zh.xlsx' and index == 2:
            wanted[5] = '=HYPERLINK("https://example.invalid", "literal text")'
        assert got == wanted, (name, index, got, wanted)
        assert sheet.cell(index, 5).data_type == 's'
        for column in [7, 8, 9]:
            if sheet.cell(index, column).value != '':
                assert sheet.cell(index, column).data_type == 'n'
    assert workbook['Files'].max_row == 3
    assert not [c for sh in workbook for row in sh for c in row if c.data_type == 'f']
    checks.append({'file': name, 'transactions': 8, 'fieldChecks': 72,
                   'formulas': 0, 'nativeAmountCells': 16, 'textDates': 8})
(root/'test-results/workbook-verified.json').write_text(json.dumps(checks, indent=2)+'\n')
print(json.dumps(checks))

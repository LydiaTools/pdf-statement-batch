from openpyxl import load_workbook

for mode in ('text', 'numeric'):
    workbook = load_workbook(f'test-results/export-{mode}.xlsx')
    sheet = workbook.worksheets[0]
    actual = [(sheet.cell(row, 7).value, sheet.cell(row, 7).data_type) for row in range(2, 6)]
    expected = ([(' 1,234.56 ', 's'), (' -12.30 ', 's'), ('0.00', 's'), ('1234567890123456.78', 's')]
                if mode == 'text' else [(1234.56, 'n'), (-12.3, 'n'), (0, 'n'), ('1234567890123456.78', 's')])
    assert actual == expected, (mode, actual)
    assert sheet['F2'].value == '=1+1' and sheet['F2'].data_type == 's'
print('Text and numeric Excel downloads preserve expected values and cell types.')
